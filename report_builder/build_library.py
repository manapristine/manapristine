"""
Mana Pristine Library Data Processor

Converts the Excel workbook 'db/library/library-books.xlsx' into a structured JSON dataset
in 'docs/library-data.json' for the static GitHub Pages library portal.

Enforces Excel as the single source of truth (openpyxl data_only=True) and validates
formula totals against individual book copy counts.
"""

from __future__ import annotations

import argparse
import json
import logging
import re
import sys
import warnings
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl

# Filter openpyxl extension warnings for clean CLI output
warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
logger = logging.getLogger(__name__)

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_EXCEL_PATH = PROJECT_ROOT / "db" / "library" / "library-books.xlsx"
DEFAULT_OUTPUT_PATH = PROJECT_ROOT / "docs" / "library-data.json"

# Standard display mapping for genres
GENRE_DISPLAY_MAP: Dict[str, str] = {
    "gre": "GRE",
    "new-age": "New Age",
    "self-help": "Self-Help",
    "computer science": "Computer Science",
    "general studies": "General Studies",
    "political science": "Political Science",
    "ayurveda": "Ayurveda",
}


def format_genre_display(genre: str) -> str:
    """Format a raw genre key into a clean display title."""
    clean = genre.strip()
    key = clean.lower()
    if key in GENRE_DISPLAY_MAP:
        return GENRE_DISPLAY_MAP[key]
    # Handle hyphenated terms like 'self-help'
    parts = clean.split("-")
    if len(parts) > 1:
        return "-".join(p.capitalize() for p in parts)
    # Default title casing
    return clean.title()


def format_language_display(lang: str) -> str:
    """Normalize language capitalization (e.g. 'hindi' -> 'Hindi')."""
    clean = lang.strip()
    return clean.capitalize()


def extract_library_data(excel_path: Path) -> Dict[str, Any]:
    """
    Parse library books and reference genres from the Excel workbook.

    Adheres to the Excel Single Source of Truth policy:
    - Loads with data_only=True to read evaluated formula values.
    - Validates the total cached formula cell (E96 / 'TOTAL BOOKS').
    """
    if not excel_path.exists():
        raise FileNotFoundError(f"Library Excel file not found: {excel_path}")

    logger.info(f"Loading workbook: {excel_path}")
    wb = openpyxl.load_workbook(excel_path, data_only=True)

    if "book" not in wb.sheetnames:
        raise ValueError(f"Required sheet 'book' missing from {excel_path.name}")

    ws_book = wb["book"]

    # Read reference genres and languages if 'genre' sheet exists
    reference_genres: List[str] = []
    reference_languages: List[str] = []
    if "genre" in wb.sheetnames:
        ws_genre = wb["genre"]
        for row in ws_genre.iter_rows(values_only=True):
            g_val = row[0]
            l_val = row[1] if len(row) > 1 else None
            if g_val and isinstance(g_val, str) and g_val.strip().upper() != "TOTAL BOOKS":
                reference_genres.append(g_val.strip())
            if l_val and isinstance(l_val, str) and l_val.strip():
                reference_languages.append(l_val.strip().capitalize())

    # Read book rows
    rows = list(ws_book.iter_rows(values_only=True))
    if not rows:
        raise ValueError("The 'book' sheet is empty.")

    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    expected_cols = ["Sr#", "Book Name", "Author", "Genre", "No. of Books", "Language", "Donated by", "Flat#"]
    for col in expected_cols:
        if col not in header:
            raise ValueError(f"Expected column '{col}' missing from header: {header}")

    col_idx = {name: header.index(name) for name in expected_cols}

    books: List[Dict[str, Any]] = []
    excel_total_copies: Optional[int] = None
    calculated_copies = 0

    for row_num, row in enumerate(rows[1:], start=2):
        sr_val = row[col_idx["Sr#"]]
        genre_val = row[col_idx["Genre"]]
        copies_val = row[col_idx["No. of Books"]]

        # Detect total row (e.g. Genre = 'TOTAL BOOKS' or Sr# is None)
        if (isinstance(genre_val, str) and genre_val.strip().upper() == "TOTAL BOOKS") or (sr_val is None and copies_val is not None):
            if copies_val is None:
                raise ValueError(
                    f"Row {row_num}: Cached formula value for 'TOTAL BOOKS' is missing (None). "
                    "Please open the workbook in Excel, save it, and close it before running this script."
                )
            try:
                excel_total_copies = int(copies_val)
            except (ValueError, TypeError) as e:
                raise ValueError(f"Row {row_num}: Invalid total copies value: {copies_val}") from e
            continue

        # Ignore empty rows if any
        if sr_val is None and row[col_idx["Book Name"]] is None:
            continue

        # Validate book row
        if not isinstance(sr_val, int):
            try:
                sr_id = int(sr_val)
            except (ValueError, TypeError) as e:
                raise ValueError(f"Row {row_num}: Invalid Sr# '{sr_val}'") from e
        else:
            sr_id = sr_val

        title = str(row[col_idx["Book Name"]]).strip() if row[col_idx["Book Name"]] is not None else ""
        author = str(row[col_idx["Author"]]).strip() if row[col_idx["Author"]] is not None else "Unknown"
        raw_genre = str(genre_val).strip() if genre_val is not None else "General"
        copies = int(copies_val) if copies_val is not None else 1
        raw_lang = str(row[col_idx["Language"]]).strip() if row[col_idx["Language"]] is not None else "English"
        donor = str(row[col_idx["Donated by"]]).strip() if row[col_idx["Donated by"]] is not None else "Community"
        flat = str(row[col_idx["Flat#"]]).strip().upper() if row[col_idx["Flat#"]] is not None else "SOCIETY"

        if not title:
            raise ValueError(f"Row {row_num}: Missing Book Name")

        calculated_copies += copies

        books.append({
            "id": sr_id,
            "title": title,
            "author": author,
            "genre": raw_genre.lower(),
            "genre_display": format_genre_display(raw_genre),
            "copies": copies,
            "language": format_language_display(raw_lang),
            "donated_by": donor,
            "flat": flat,
        })

    # Validate against Excel formula total
    if excel_total_copies is not None:
        if calculated_copies != excel_total_copies:
            raise ValueError(
                f"Mismatch in total book copies! Excel cached formula evaluates to {excel_total_copies}, "
                f"but sum of rows is {calculated_copies}."
            )
        logger.info(f"Verified formula total: {excel_total_copies} total copies across {len(books)} titles.")
    else:
        logger.warning("No 'TOTAL BOOKS' formula row detected; using calculated copies count.")
        excel_total_copies = calculated_copies

    # Build aggregations
    genre_stats: Dict[str, Dict[str, Any]] = {}
    lang_stats: Dict[str, Dict[str, Any]] = {}
    donor_stats: Dict[str, Dict[str, Any]] = {}

    for b in books:
        # Genre breakdown
        g_key = b["genre"]
        if g_key not in genre_stats:
            genre_stats[g_key] = {
                "genre": g_key,
                "display": b["genre_display"],
                "titles": 0,
                "copies": 0,
            }
        genre_stats[g_key]["titles"] += 1
        genre_stats[g_key]["copies"] += b["copies"]

        # Language breakdown
        l_key = b["language"]
        if l_key not in lang_stats:
            lang_stats[l_key] = {
                "language": l_key,
                "titles": 0,
                "copies": 0,
            }
        lang_stats[l_key]["titles"] += 1
        lang_stats[l_key]["copies"] += b["copies"]

        # Donor breakdown
        d_key = b["donated_by"]
        if d_key not in donor_stats:
            donor_stats[d_key] = {
                "donor": d_key,
                "flat": b["flat"],
                "titles": 0,
                "copies": 0,
            }
        donor_stats[d_key]["titles"] += 1
        donor_stats[d_key]["copies"] += b["copies"]

    sorted_genres = sorted(genre_stats.values(), key=lambda x: (-x["titles"], x["display"]))
    sorted_langs = sorted(lang_stats.values(), key=lambda x: (-x["copies"], x["language"]))
    sorted_donors = sorted(donor_stats.values(), key=lambda x: (-x["copies"], x["donor"]))

    payload = {
        "metadata": {
            "title": "Mana Pristine Community Library",
            "source_file": str(excel_path.name),
            "generated_at": datetime.now().isoformat(),
            "version": "1.0",
        },
        "summary": {
            "total_titles": len(books),
            "total_copies": calculated_copies,
            "excel_total_copies": excel_total_copies,
            "unique_genres": len(genre_stats),
            "unique_languages": len(lang_stats),
            "unique_donors": len(donor_stats),
            "genre_breakdown": sorted_genres,
            "language_breakdown": sorted_langs,
            "donor_breakdown": sorted_donors,
        },
        "reference": {
            "genres": reference_genres,
            "languages": reference_languages,
        },
        "books": books,
    }

    return payload


def convert_library_excel_to_json(
    excel_path: Path, output_json_path: Path, dry_run: bool = False
) -> Dict[str, Any]:
    """Convert the library Excel workbook to JSON and write to disk."""
    data = extract_library_data(excel_path)

    if dry_run:
        logger.info(f"[Dry Run] Validated {len(data['books'])} books. Output skipped.")
        return data

    output_json_path.parent.mkdir(parents=True, exist_ok=True)
    with open(output_json_path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)

    logger.info(f"Successfully generated {output_json_path} ({len(data['books'])} titles, {data['summary']['total_copies']} copies)")
    return data


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Convert Mana Pristine library-books.xlsx to JSON for GitHub Pages portal."
    )
    parser.add_argument(
        "--excel",
        type=Path,
        default=DEFAULT_EXCEL_PATH,
        help="Path to library-books.xlsx workbook (default: db/library/library-books.xlsx)",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT_PATH,
        help="Path to output JSON file (default: docs/library-data.json)",
    )
    parser.add_argument(
        "--check",
        action="store_true",
        help="Validate Excel data without writing JSON",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    try:
        convert_library_excel_to_json(args.excel, args.output, dry_run=args.check)
        return 0
    except Exception as e:
        logger.error(f"Failed to process library data: {e}")
        return 1


if __name__ == "__main__":
    sys.exit(main())
