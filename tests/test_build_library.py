import json
import tempfile
import unittest
from pathlib import Path
import openpyxl

from report_builder.build_library import (
    format_genre_display,
    format_language_display,
    extract_library_data,
    convert_library_excel_to_json,
)


class TestBuildLibrary(unittest.TestCase):

    def test_format_genre_display(self):
        self.assertEqual(format_genre_display("fiction"), "Fiction")
        self.assertEqual(format_genre_display("gre"), "GRE")
        self.assertEqual(format_genre_display("new-age"), "New Age")
        self.assertEqual(format_genre_display("self-help"), "Self-Help")
        self.assertEqual(format_genre_display("computer science"), "Computer Science")
        self.assertEqual(format_genre_display("  philosophy  "), "Philosophy")

    def test_format_language_display(self):
        self.assertEqual(format_language_display("hindi"), "Hindi")
        self.assertEqual(format_language_display("English"), "English")
        self.assertEqual(format_language_display("bangla"), "Bangla")
        self.assertEqual(format_language_display("  tamil "), "Tamil")

    def test_extract_actual_library_workbook(self):
        excel_path = Path("db/library/library-books.xlsx")
        if not excel_path.exists():
            self.skipTest(f"Skipping: {excel_path} does not exist")

        data = extract_library_data(excel_path)
        self.assertIn("metadata", data)
        self.assertIn("summary", data)
        self.assertIn("books", data)

        summary = data["summary"]
        self.assertEqual(summary["total_titles"], 94)
        self.assertEqual(summary["total_copies"], 147)
        self.assertEqual(summary["excel_total_copies"], 147)
        self.assertGreater(summary["unique_genres"], 20)
        self.assertGreaterEqual(summary["unique_languages"], 3)

        # Check first book
        b0 = data["books"][0]
        self.assertEqual(b0["id"], 1)
        self.assertEqual(b0["title"], "Crossing the Chasm")
        self.assertEqual(b0["author"], "Geoffrey Moore")
        self.assertEqual(b0["genre"], "marketing")
        self.assertEqual(b0["genre_display"], "Marketing")
        self.assertEqual(b0["copies"], 1)
        self.assertEqual(b0["language"], "English")
        self.assertEqual(b0["donated_by"], "Samrat Kar")
        self.assertEqual(b0["flat"], "F3")

    def test_custom_workbook_processing_and_conversion(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            wb_path = Path(tmpdir) / "test-library.xlsx"
            out_json = Path(tmpdir) / "test-out.json"

            wb = openpyxl.Workbook()
            ws_book = wb.active
            ws_book.title = "book"
            ws_book.append(["Sr#", "Book Name", "Author", "Genre", "No. of Books", "Language", "Donated by", "Flat#"])
            ws_book.append([1, "Test Book 1", "Author A", "fiction", 2, "English", "Resident 1", "S10"])
            ws_book.append([2, "Test Book 2", "Author B", "science", 3, "hindi", "Resident 2", "T12"])
            ws_book.append([None, None, None, "TOTAL BOOKS", 5, None, None, None])

            ws_genre = wb.create_sheet("genre")
            ws_genre.append(["fiction", "English"])
            ws_genre.append(["science", "Hindi"])

            wb.save(wb_path)

            data = convert_library_excel_to_json(wb_path, out_json)
            self.assertTrue(out_json.exists())

            with open(out_json, "r", encoding="utf-8") as f:
                loaded = json.load(f)

            self.assertEqual(loaded["summary"]["total_titles"], 2)
            self.assertEqual(loaded["summary"]["total_copies"], 5)
            self.assertEqual(loaded["summary"]["excel_total_copies"], 5)
            self.assertEqual(len(loaded["books"]), 2)
            self.assertEqual(loaded["books"][1]["language"], "Hindi")
            self.assertEqual(loaded["books"][0]["flat"], "S10")

    def test_missing_sheet_raises_error(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            wb_path = Path(tmpdir) / "empty.xlsx"
            wb = openpyxl.Workbook()
            wb.save(wb_path)

            with self.assertRaises(ValueError) as ctx:
                extract_library_data(wb_path)
            self.assertIn("Required sheet 'book' missing", str(ctx.exception))

    def test_formula_mismatch_raises_error(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            wb_path = Path(tmpdir) / "mismatch.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "book"
            ws.append(["Sr#", "Book Name", "Author", "Genre", "No. of Books", "Language", "Donated by", "Flat#"])
            ws.append([1, "Book 1", "Author 1", "fiction", 1, "English", "Donor", "F1"])
            # Say total is 99 instead of 1
            ws.append([None, None, None, "TOTAL BOOKS", 99, None, None, None])
            wb.save(wb_path)

            with self.assertRaises(ValueError) as ctx:
                extract_library_data(wb_path)
            self.assertIn("Mismatch in total book copies", str(ctx.exception))


if __name__ == "__main__":
    unittest.main()
