import unittest
import tempfile
import shutil
from pathlib import Path
import openpyxl

from report_builder.track_payments import (
    create_tracking_workbook,
    create_unmatched_report,
    find_flat_in_description,
    track_payments,
)


class TestTrackPayments(unittest.TestCase):

    def setUp(self):
        self.test_dir = Path(tempfile.mkdtemp())
        self.mappings = {
            "flat_to_names": {
                "f1": {"shailesh", "srivastava", "shaileshsrivastava", "nikhil", "nikhilk"},
                "g2": {"anand", "kumar", "anandkumar"},
            },
            "name_part_to_flats": {
                "shailesh": {"f1"},
                "srivastava": {"f1"},
                "nikhil": {"f1"},
                "anand": {"g2"},
                "kumar": {"g2"},
            },
            "flat_to_name_parts": {
                "f1": [["shailesh", "srivastava"], ["nikhil", "k"]],
                "g2": [["anand", "kumar"]],
            },
            "flat_to_member": {"f1": "Shailesh Srivastava", "g2": "Anand Kumar"},
            "flat_to_occupant": {"f1": "Nikhil K", "g2": ""},
            "owner_names": {"f1": {"shaileshsrivastava"}, "g2": {"anandkumar"}},
            "occupant_names": {"f1": {"nikhilk", "nikhil"}},
            "all_flats": ["F1", "G2"],
        }

    def tearDown(self):
        shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_find_flat_in_description(self):
        # Direct flat match
        flat, reason = find_flat_in_description("Maintenance payment for Flat F1", self.mappings)
        self.assertEqual(flat, "f1")
        self.assertEqual(reason, "Flat # in description")

        # Owner name match
        flat, reason = find_flat_in_description("NEFT from Shailesh Srivastava", self.mappings)
        self.assertEqual(flat, "f1")
        self.assertEqual(reason, "Owner Name match")

        # Occupant name match
        flat, reason = find_flat_in_description("UPI transfer Nikhil K", self.mappings)
        self.assertEqual(flat, "f1")
        self.assertEqual(reason, "Occupant Name match")

        # Unmatched
        flat, reason = find_flat_in_description("Interest Credit SBI", self.mappings)
        self.assertIsNone(flat)
        self.assertEqual(reason, "Unmatched")

    def test_create_tracking_workbook_structure(self):
        transactions = [
            {
                "date": "05-Aug-2026",
                "value_date": "05-Aug-2026",
                "amount": 3500.0,
                "flat": "F1",
                "owner": "Shailesh Srivastava",
                "occupant": "Nikhil K",
                "match_reason": "Flat # in description",
                "description": "UPI/123456789/MAINT F1",
                "ref_no": "123456789",
            },
            {
                "date": "10-Aug-2026",
                "value_date": "10-Aug-2026",
                "amount": 500.0,
                "flat": "F1",
                "owner": "Shailesh Srivastava",
                "occupant": "Nikhil K",
                "match_reason": "Flat # in description",
                "description": "UPI/987654321/FINE F1",
                "ref_no": "987654321",
            },
            {
                "date": "06-Aug-2026",
                "value_date": "06-Aug-2026",
                "amount": 250.0,
                "flat": None,
                "owner": "",
                "occupant": "",
                "match_reason": "Unmatched",
                "description": "BANK CHARGES REVERSAL",
                "ref_no": "REV001",
            },
        ]

        out_path = self.test_dir / "test_tracking.xlsx"
        create_tracking_workbook(transactions, self.mappings, out_path, statement_name="sample_statement.xls")

        self.assertTrue(out_path.exists())
        wb = openpyxl.load_workbook(out_path)

        # Verify sheets
        self.assertIn("Payment Tracking", wb.sheetnames)
        self.assertIn("All Transactions", wb.sheetnames)
        self.assertIn("Unmatched Transactions", wb.sheetnames)

        # Verify Payment Tracking headers (18 cols with max 3 payments, 4 cols each)
        ws_track = wb["Payment Tracking"]
        headers = [c.value for c in ws_track[4]]
        self.assertEqual(headers[0], "Flat #")
        self.assertEqual(headers[1], "Owner Name (Member)")
        self.assertEqual(headers[2], "Occupant Name (Tenant)")
        self.assertEqual(headers[3], "Status")
        self.assertEqual(headers[4], "Total Paid (Rs.)")
        self.assertEqual(headers[5], "Payments Count")
        self.assertEqual(headers[6], "Payment 1 Amount (Rs.)")
        self.assertEqual(headers[7], "Payment 1 Date")
        self.assertEqual(headers[8], "Payment 1 Ref No. / Cheque No.")
        self.assertEqual(headers[9], "Payment 1 Bank Transaction Details")
        self.assertEqual(headers[10], "Payment 2 Amount (Rs.)")
        self.assertEqual(headers[11], "Payment 2 Date")
        self.assertEqual(headers[12], "Payment 2 Ref No. / Cheque No.")
        self.assertEqual(headers[13], "Payment 2 Bank Transaction Details")

        # Verify row 5 (F1 has 2 payments)
        row_f1 = [c.value for c in ws_track[5]]
        self.assertEqual(row_f1[0], "F1")
        self.assertEqual(row_f1[3], "PAID")
        self.assertEqual(row_f1[4], 4000.0)
        self.assertEqual(row_f1[5], 2)
        self.assertEqual(row_f1[6], 3500.0)
        self.assertEqual(row_f1[7], "05-Aug-2026")
        self.assertEqual(row_f1[8], "123456789")
        self.assertEqual(row_f1[9], "UPI/123456789/MAINT F1")
        self.assertEqual(row_f1[10], 500.0)
        self.assertEqual(row_f1[11], "10-Aug-2026")
        self.assertEqual(row_f1[12], "987654321")
        self.assertEqual(row_f1[13], "UPI/987654321/FINE F1")

        # Verify row 6 (G2 has 0 payments - PENDING)
        row_g2 = [c.value for c in ws_track[6]]
        self.assertEqual(row_g2[0], "G2")
        self.assertEqual(row_g2[3], "PENDING")
        self.assertEqual(row_g2[4], 0.0)
        self.assertEqual(row_g2[5], 0)
        self.assertIn(row_g2[6], [None, ""])

        # Verify All Transactions
        ws_tx = wb["All Transactions"]
        self.assertEqual(ws_tx["K4"].value, "Bank Transaction Details / Description")
        self.assertEqual(ws_tx["K5"].value, "UPI/123456789/MAINT F1")

    def test_create_unmatched_report_xlsx(self):
        unmatched = [
            {
                "date": "06-Aug-2026",
                "value_date": "06-Aug-2026",
                "amount": 250.0,
                "flat": None,
                "match_reason": "Unmatched",
                "description": "BANK CHARGES REVERSAL",
                "ref_no": "REV001",
            }
        ]
        out_path = self.test_dir / "unmatched.xlsx"
        create_unmatched_report(unmatched, out_path, statement_name="stmt.xls")

        self.assertTrue(out_path.exists())
        wb = openpyxl.load_workbook(out_path)
        ws = wb.active
        self.assertEqual(ws["G5"].value, "Bank Transaction Details / Description")
        self.assertEqual(ws["H5"].value, "Manual Flat #")
        self.assertEqual(ws["G6"].value, "BANK CHARGES REVERSAL")


if __name__ == "__main__":
    unittest.main()
